# -*- coding: utf-8 -*-
"""Geração do relatório de apontamentos em DOCX e PDF, com a identidade
visual do escritório (preto e dourado, logo Zera, títulos em Cinzel).

Os dois formatos saem do mesmo conteúdo (montar_conteudo), então o que o
cliente lê no PDF é o mesmo que está no Word.
"""
from __future__ import annotations

import re
import unicodedata
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from .base import MESES, PRIORIDADES, competencia_extenso

MARCA = Path(__file__).resolve().parent / "marca"
LOGO = MARCA / "logo_claro.png"

OURO = (168, 132, 63)
OURO_CLARO = (201, 166, 98)
OURO_SUAVE = (246, 240, 228)
PRETO = (17, 17, 17)
CINZA = (107, 102, 92)
CINZA_LINHA = (228, 222, 210)
COR_PRIORIDADE = {"Alta": (196, 56, 61), "Média": (168, 132, 63), "Baixa": (31, 138, 91)}

TITULO_DOC = "Relatório de Apontamentos Contábeis"


class ErroDocumento(Exception):
    pass


# ----------------------------------------------------------------------
# conteúdo comum aos dois formatos
def data_extenso(d: date) -> str:
    return f"{d.day:02d} de {MESES[d.month - 1].lower()} de {d.year}"


def formatar_valor(valor: str) -> str:
    v = (valor or "").strip()
    if not v:
        return ""
    bruto = re.sub(r"^R\$\s*", "", v, flags=re.I).strip()
    numero = None
    try:
        if re.fullmatch(r"\d{1,3}(\.\d{3})+(,\d{1,2})?|\d+(,\d{1,2})?", bruto):
            numero = Decimal(bruto.replace(".", "").replace(",", "."))
        elif re.fullmatch(r"\d+\.\d{1,2}", bruto):
            numero = Decimal(bruto)
    except InvalidOperation:
        numero = None
    if numero is None:
        return v if v.upper().startswith("R$") else f"R$ {v}" if re.match(r"^[\d.,]+$", v) else v
    inteiro, centavos = f"{numero:.2f}".split(".")
    inteiro = f"{int(inteiro):,}".replace(",", ".")
    return f"R$ {inteiro},{centavos}"


def paragrafos(texto: str) -> list[str]:
    partes = re.split(r"\n\s*\n", (texto or "").strip())
    return [re.sub(r"\s*\n\s*", " ", p).strip() for p in partes if p.strip()]


def montar_conteudo(doc: dict, cfg, hoje: date | None = None) -> dict:
    hoje = hoje or date.today()
    comp = competencia_extenso(doc["competencia"])
    itens = []
    for n, a in enumerate(doc.get("apontamentos", []), 1):
        titulo = a.get("titulo") or _titulo_reserva(a)
        itens.append({
            "id": a.get("id"),
            "situacao": a.get("situacao", ""),
            "numero": n,
            "titulo": titulo,
            "categoria": a.get("categoria", ""),
            "prioridade": a.get("prioridade", "Média"),
            "paragrafos": paragrafos(a.get("texto") or a.get("original", "")),
            "providencia": (a.get("providencia") or "").strip(),
            "valor": formatar_valor(a.get("valor", "")),
            "referencia": (a.get("referencia") or "").strip(),
        })
    por_prioridade = {p: sum(1 for i in itens if i["prioridade"] == p) for p in PRIORIDADES}
    por_categoria: dict = {}
    for i in itens:
        por_categoria[i["categoria"]] = por_categoria.get(i["categoria"], 0) + 1

    prazo = ""
    if doc.get("prazo_retorno"):
        try:
            prazo = datetime.strptime(doc["prazo_retorno"], "%Y-%m-%d").strftime("%d/%m/%Y")
        except ValueError:
            prazo = ""

    def _subst(t: str) -> str:
        return (t or "").replace("{competencia}", comp).replace(
            "{empresa}", doc["empresa"]).replace("{prazo}", prazo or "")

    encerramento = _subst(cfg.encerramento)
    if prazo and "{prazo}" not in (cfg.encerramento or ""):
        encerramento = (f"Solicitamos o retorno até {prazo}. " + encerramento).strip()

    rodape = [x for x in (cfg.escritorio_telefone, cfg.escritorio_email,
                          cfg.escritorio_site) if x]
    return {
        "titulo": TITULO_DOC,
        "empresa": doc["empresa"],
        "cnpj": doc.get("cnpj", ""),
        "competencia": comp,
        "destinatario": doc.get("destinatario", ""),
        "prazo": prazo,
        "emissao": hoje.strftime("%d/%m/%Y"),
        "data_extenso": data_extenso(hoje),
        "introducao": paragrafos(_subst(cfg.introducao)),
        "encerramento": paragrafos(encerramento),
        "itens": itens,
        "por_prioridade": por_prioridade,
        "por_categoria": por_categoria,
        "escritorio": cfg.escritorio_nome or "Zera Contabilidade",
        "escritorio_cnpj": cfg.escritorio_cnpj,
        "escritorio_endereco": cfg.escritorio_endereco,
        "rodape": " · ".join(rodape),
        "responsavel": cfg.responsavel_nome,
        "cargo": cfg.responsavel_cargo,
        "crc": cfg.responsavel_crc,
    }


def _titulo_reserva(a: dict) -> str:
    return a.get("categoria") or "Apontamento"


def nome_arquivo(doc: dict) -> str:
    empresa = unicodedata.normalize("NFKC", doc["empresa"])
    empresa = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "", empresa).strip(" .")[:80] or "Empresa"
    ano, mes = doc["competencia"].split("-")
    return f"Apontamentos - {empresa} - {mes}-{ano}"


def pasta_destino(cfg, doc: dict) -> Path:
    empresa = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "", doc["empresa"]).strip(" .")[:80] or "Empresa"
    p = Path(cfg.pasta_saida) / empresa
    p.mkdir(parents=True, exist_ok=True)
    return p


def gerar(doc: dict, cfg, formatos=("docx", "pdf"), pasta: Path | None = None) -> list[Path]:
    if not doc.get("apontamentos"):
        raise ErroDocumento("O documento não tem nenhum apontamento.")
    pendentes = [a for a in doc["apontamentos"] if a.get("situacao") == "pendente"]
    if pendentes:
        raise ErroDocumento("Ainda há apontamento sendo formalizado pela IA. "
                            "Aguarde terminar e tente de novo.")
    c = montar_conteudo(doc, cfg)
    pasta = pasta or pasta_destino(cfg, doc)
    base = pasta / nome_arquivo(doc)
    saidas = []
    for fmt in formatos:
        destino = base.with_suffix("." + fmt)
        try:
            if fmt == "docx":
                gerar_docx(c, destino)
            elif fmt == "pdf":
                gerar_pdf(c, destino)
            else:
                continue
        except PermissionError as e:
            raise ErroDocumento(f"Não consegui gravar {destino.name}. Se ele estiver "
                                f"aberto no Word ou no leitor de PDF, feche e tente de novo.") from e
        saidas.append(destino)
    return saidas


# ----------------------------------------------------------------------
# PDF (fpdf2)
def gerar_pdf(c: dict, destino: Path) -> None:
    from fpdf import FPDF
    from fpdf.enums import XPos, YPos

    class PDF(FPDF):
        def header(self):
            if LOGO.exists():
                self.image(str(LOGO), x=18, y=11, h=13)
            self.set_xy(80, 12)
            self.set_font("Cinzel", "B", 9)
            self.set_text_color(*OURO)
            self.cell(112, 5, "APONTAMENTOS CONTÁBEIS", align="R",
                      new_x=XPos.LEFT, new_y=YPos.NEXT)
            self.set_font("Corpo", "", 8.5)
            self.set_text_color(*CINZA)
            self.cell(112, 5, f"{c['empresa']} · {c['competencia']}"[:90], align="R")
            self.set_draw_color(*OURO)
            self.set_line_width(0.6)
            self.line(18, 28, 192, 28)
            self.set_line_width(0.2)
            self.set_y(34)

        def footer(self):
            self.set_y(-17)
            self.set_draw_color(*OURO_CLARO)
            self.set_line_width(0.3)
            self.line(18, self.get_y(), 192, self.get_y())
            self.ln(2)
            self.set_font("Corpo", "", 7.5)
            self.set_text_color(*CINZA)
            esquerda = c["escritorio"] + (f" · {c['rodape']}" if c["rodape"] else "")
            self.cell(140, 4, esquerda[:120])
            self.cell(34, 4, f"Página {self.page_no()} de {{nb}}", align="R")

    pdf = PDF(format="A4")
    pdf.set_margins(18, 34, 18)
    pdf.set_auto_page_break(True, margin=22)
    pdf.add_font("Cinzel", "B", str(MARCA / "Cinzel-Bold.ttf"))
    pdf.add_font("Cinzel", "", str(MARCA / "Cinzel-Medium.ttf"))
    pdf.add_font("Corpo", "", str(MARCA / "LiberationSans-Regular.ttf"))
    pdf.add_font("Corpo", "B", str(MARCA / "LiberationSans-Bold.ttf"))
    pdf.add_font("Corpo", "I", str(MARCA / "LiberationSans-Italic.ttf"))
    pdf.add_font("Corpo", "BI", str(MARCA / "LiberationSans-BoldItalic.ttf"))
    pdf.set_title(f"{c['titulo']} - {c['empresa']} - {c['competencia']}")
    pdf.set_author(c["escritorio"])
    pdf.set_creator("Apontamentos Contábeis IA")
    pdf.alias_nb_pages()
    pdf.add_page()
    largura = 174

    # título
    pdf.set_font("Cinzel", "B", 18)
    pdf.set_text_color(*PRETO)
    pdf.multi_cell(largura, 9, c["titulo"], new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.set_font("Corpo", "", 10)
    pdf.set_text_color(*CINZA)
    pdf.cell(largura, 6, f"Competência {c['competencia']}", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.ln(4)

    # quadro de identificação
    linhas = [("Cliente", c["empresa"])]
    if c["cnpj"]:
        linhas.append(("CNPJ", c["cnpj"]))
    linhas.append(("Competência", c["competencia"]))
    if c["destinatario"]:
        linhas.append(("A/C", c["destinatario"]))
    linhas.append(("Emissão", c["emissao"]))
    if c["prazo"]:
        linhas.append(("Retorno até", c["prazo"]))
    y0 = pdf.get_y()
    altura = 6.2 * len(linhas) + 6
    pdf.set_fill_color(*OURO_SUAVE)
    pdf.rect(18, y0, largura, altura, style="F")
    pdf.set_fill_color(*OURO)
    pdf.rect(18, y0, 1.4, altura, style="F")
    pdf.set_y(y0 + 3)
    for rot, val in linhas:
        pdf.set_x(24)
        pdf.set_font("Corpo", "B", 9.5)
        pdf.set_text_color(*CINZA)
        pdf.cell(32, 6.2, rot.upper())
        pdf.set_font("Corpo", "", 10)
        pdf.set_text_color(*PRETO)
        pdf.cell(largura - 40, 6.2, val[:95], new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.set_y(y0 + altura + 7)

    def texto_corrido(t: str, tamanho=10.5, estilo=""):
        pdf.set_font("Corpo", estilo, tamanho)
        pdf.set_text_color(*PRETO)
        pdf.multi_cell(largura, 5.6, t, align="J", new_x=XPos.LMARGIN, new_y=YPos.NEXT)

    def subtitulo(t: str):
        if pdf.get_y() > 250:
            pdf.add_page()
        pdf.set_font("Cinzel", "B", 11.5)
        pdf.set_text_color(*OURO)
        pdf.cell(largura, 7, t.upper(), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        pdf.set_draw_color(*CINZA_LINHA)
        pdf.line(18, pdf.get_y(), 192, pdf.get_y())
        pdf.ln(3)

    pdf.set_font("Corpo", "", 10.5)
    pdf.set_text_color(*PRETO)
    pdf.cell(largura, 6, "Prezados(as),", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.ln(1)
    for p in c["introducao"]:
        texto_corrido(p)
        pdf.ln(2)
    pdf.ln(2)

    # resumo
    subtitulo("Resumo")
    total = len(c["itens"])
    pdf.set_font("Corpo", "", 10)
    pdf.set_text_color(*PRETO)
    partes = [f"{total} apontamento{'s' if total != 1 else ''}"]
    for p in PRIORIDADES:
        if c["por_prioridade"].get(p):
            partes.append(f"{c['por_prioridade'][p]} de prioridade {p.lower()}")
    pdf.multi_cell(largura, 5.6, " · ".join(partes), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.ln(1.5)
    for cat, n in c["por_categoria"].items():
        pdf.set_x(22)
        pdf.set_text_color(*OURO)
        pdf.cell(4, 5.4, "▪")
        pdf.set_text_color(*PRETO)
        pdf.cell(130, 5.4, cat)
        pdf.set_text_color(*CINZA)
        pdf.cell(20, 5.4, str(n), align="R", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.ln(5)

    # apontamentos
    subtitulo("Apontamentos")
    for it in c["itens"]:
        if pdf.get_y() > 248:          # título sem o texto logo abaixo: nova página
            pdf.add_page()
        pdf.ln(1)
        pdf.set_font("Cinzel", "B", 13)
        pdf.set_text_color(*OURO)
        pdf.cell(12, 7, f"{it['numero']:02d}")
        pdf.set_font("Corpo", "B", 11.5)
        pdf.set_text_color(*PRETO)
        pdf.multi_cell(largura - 12, 7, it["titulo"], new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        pdf.set_x(30)
        pdf.set_font("Corpo", "", 8.5)
        pdf.set_text_color(*CINZA)
        pdf.cell(pdf.get_string_width(it["categoria"]) + 3, 5, it["categoria"])
        cor = COR_PRIORIDADE.get(it["prioridade"], OURO)
        pdf.set_text_color(*cor)
        pdf.set_font("Corpo", "B", 8.5)
        pdf.cell(60, 5, f"●  Prioridade {it['prioridade'].lower()}",
                 new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        pdf.ln(1.5)
        for p in it["paragrafos"]:
            pdf.set_x(30)
            pdf.set_font("Corpo", "", 10.5)
            pdf.set_text_color(*PRETO)
            pdf.multi_cell(largura - 12, 5.6, p, align="J", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
            pdf.ln(1.2)
        detalhes = []
        if it["referencia"]:
            detalhes.append(("Referência", it["referencia"]))
        if it["valor"]:
            detalhes.append(("Valor", it["valor"]))
        for rot, val in detalhes:
            pdf.set_x(30)
            pdf.set_font("Corpo", "B", 9.5)
            pdf.set_text_color(*CINZA)
            pdf.cell(22, 5.4, rot + ":")
            pdf.set_font("Corpo", "", 9.5)
            pdf.set_text_color(*PRETO)
            pdf.multi_cell(largura - 34, 5.4, val, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        if it["providencia"]:
            pdf.ln(1)
            pdf.set_font("Corpo", "", 10)
            linhas_prov = pdf.multi_cell(largura - 18, 5.4, it["providencia"],
                                         dry_run=True, output="LINES")
            h = 5.4 * len(linhas_prov) + 10
            if pdf.get_y() + h > 275:
                pdf.add_page()
            yb = pdf.get_y()
            pdf.set_fill_color(*OURO_SUAVE)
            pdf.rect(30, yb, largura - 12, h, style="F")
            pdf.set_fill_color(*OURO)
            pdf.rect(30, yb, 1.2, h, style="F")
            pdf.set_xy(34, yb + 2)
            pdf.set_font("Corpo", "B", 8.5)
            pdf.set_text_color(*OURO)
            pdf.cell(80, 4.5, "PROVIDÊNCIA SOLICITADA", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
            pdf.set_x(34)
            pdf.set_font("Corpo", "", 10)
            pdf.set_text_color(*PRETO)
            pdf.multi_cell(largura - 18, 5.4, it["providencia"], new_x=XPos.LMARGIN,
                           new_y=YPos.NEXT)
            pdf.set_y(yb + h)
        pdf.ln(5)

    # encerramento e assinatura
    if pdf.get_y() > 215:
        pdf.add_page()
    subtitulo("Considerações finais")
    for p in c["encerramento"]:
        texto_corrido(p)
        pdf.ln(2)
    pdf.ln(4)
    if pdf.get_y() > 228:              # assinatura não fica separada do fecho
        pdf.add_page()
    pdf.set_font("Corpo", "", 10.5)
    pdf.cell(largura, 6, "Atenciosamente,", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.ln(16)
    pdf.set_draw_color(*PRETO)
    pdf.line(18, pdf.get_y(), 98, pdf.get_y())
    pdf.ln(1.5)
    if c["responsavel"]:
        pdf.set_font("Corpo", "B", 10.5)
        pdf.cell(largura, 5.5, c["responsavel"], new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.set_font("Corpo", "", 9.5)
    pdf.set_text_color(*CINZA)
    cargo = c["cargo"] + (f" · CRC {c['crc']}" if c["crc"] else "")
    if cargo.strip():
        pdf.cell(largura, 5, cargo, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.set_font("Cinzel", "B", 10)
    pdf.set_text_color(*OURO)
    pdf.cell(largura, 6, c["escritorio"].upper(), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.set_font("Corpo", "", 8.5)
    pdf.set_text_color(*CINZA)
    extra = " · ".join(x for x in (f"CNPJ {c['escritorio_cnpj']}" if c["escritorio_cnpj"] else "",
                                   c["escritorio_endereco"]) if x)
    if extra:
        pdf.multi_cell(largura, 4.5, extra, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.set_font("Corpo", "", 9.5)
    pdf.ln(3)
    pdf.cell(largura, 5, c["data_extenso"], new_x=XPos.LMARGIN, new_y=YPos.NEXT)

    destino.parent.mkdir(parents=True, exist_ok=True)
    tmp = destino.with_name(destino.stem + ".tmp.pdf")
    pdf.output(str(tmp))
    _trocar(tmp, destino)


# ----------------------------------------------------------------------
# DOCX (python-docx)
def gerar_docx(c: dict, destino: Path) -> None:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    rgb = lambda t: RGBColor(*t)  # noqa: E731
    hexa = lambda t: "%02X%02X%02X" % t  # noqa: E731
    TITULO = "Georgia"   # serifada presente em todo Windows, próxima da Cinzel
    CORPO = "Arial"

    d = Document()
    sec = d.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(1.8)
    sec.top_margin, sec.bottom_margin = Cm(3.2), Cm(2.2)
    sec.header_distance, sec.footer_distance = Cm(1.0), Cm(1.0)

    normal = d.styles["Normal"]
    normal.font.name = CORPO
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), CORPO)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.15

    def run(p, texto, *, tam=None, negrito=False, italico=False, cor=None, fonte=None,
            caixa_alta=False, espaco=None):
        r = p.add_run(texto)
        r.bold, r.italic = negrito, italico
        if tam:
            r.font.size = Pt(tam)
        if cor:
            r.font.color.rgb = rgb(cor)
        if fonte:
            r.font.name = fonte
            r._element.rPr.rFonts.set(qn("w:eastAsia"), fonte)
        if caixa_alta:
            r.font.all_caps = True
        if espaco is not None:
            sp = OxmlElement("w:spacing")
            sp.set(qn("w:val"), str(espaco))
            r._element.get_or_add_rPr().insert_element_before(sp, *_DEPOIS_SPACING_R)
        return r

    def borda(p, lado, cor, tam=6, espaco=4):
        pPr = p._p.get_or_add_pPr()
        bdr = pPr.find(qn("w:pBdr"))
        if bdr is None:
            bdr = OxmlElement("w:pBdr")
            pPr.insert_element_before(bdr, *_DEPOIS_PBDR)
        b = OxmlElement(f"w:{lado}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), str(tam))
        b.set(qn("w:space"), str(espaco))
        b.set(qn("w:color"), hexa(cor))
        bdr.append(b)

    def sombrear(elemento_pr, cor):
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hexa(cor))
        sucessores = _DEPOIS_SHD_P if elemento_pr.tag == qn("w:pPr") else _DEPOIS_SHD_TC
        elemento_pr.insert_element_before(shd, *sucessores)

    def campo(p, codigo, tam):
        for tipo, texto in (("begin", None), (None, codigo), ("separate", None),
                            (None, "1"), ("end", None)):
            r = p.add_run()
            r.font.size = Pt(tam)
            r.font.color.rgb = rgb(CINZA)
            if tipo:
                f = OxmlElement("w:fldChar")
                f.set(qn("w:fldCharType"), tipo)
                r._element.append(f)
            elif texto == codigo:
                t = OxmlElement("w:instrText")
                t.set(qn("xml:space"), "preserve")
                t.text = f" {codigo} "
                r._element.append(t)
            else:
                r.text = texto

    # cabeçalho: logo à esquerda, identificação à direita, fio dourado
    cab = sec.header
    tab = cab.add_table(rows=1, cols=2, width=Cm(17.4))
    tab.alignment = WD_TABLE_ALIGNMENT.CENTER
    tab.autofit = False
    c0, c1 = tab.rows[0].cells
    c0.width, c1.width = Cm(6.0), Cm(11.4)
    p = c0.paragraphs[0]
    if LOGO.exists():
        p.add_run().add_picture(str(LOGO), height=Cm(1.35))
    p = c1.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run(p, "APONTAMENTOS CONTÁBEIS", tam=9, negrito=True, cor=OURO, fonte=TITULO, espaco=20)
    p = c1.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run(p, f"{c['empresa']} · {c['competencia']}", tam=8.5, cor=CINZA)
    if cab.paragraphs and not cab.paragraphs[0].text:
        primeiro = cab.paragraphs[0]
        primeiro.paragraph_format.space_after = Pt(0)
        primeiro._p.getparent().remove(primeiro._p)
    linha = cab.add_paragraph()
    linha.paragraph_format.space_after = Pt(0)
    borda(linha, "bottom", OURO, tam=12, espaco=1)

    # rodapé: escritório e paginação
    rod = sec.footer
    p = rod.paragraphs[0]
    borda(p, "top", OURO_CLARO, tam=4, espaco=4)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(17.4), alignment=2)  # direita
    esquerda = c["escritorio"] + (f" · {c['rodape']}" if c["rodape"] else "")
    run(p, esquerda, tam=7.5, cor=CINZA)
    run(p, "\tPágina ", tam=7.5, cor=CINZA)
    campo(p, "PAGE", 7.5)
    run(p, " de ", tam=7.5, cor=CINZA)
    campo(p, "NUMPAGES", 7.5)

    # título
    p = d.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    run(p, c["titulo"], tam=19, negrito=True, cor=PRETO, fonte=TITULO, caixa_alta=True)
    p = d.add_paragraph()
    run(p, f"Competência {c['competencia']}", tam=10, cor=CINZA)

    # quadro de identificação
    linhas = [("Cliente", c["empresa"])]
    if c["cnpj"]:
        linhas.append(("CNPJ", c["cnpj"]))
    linhas.append(("Competência", c["competencia"]))
    if c["destinatario"]:
        linhas.append(("A/C", c["destinatario"]))
    linhas.append(("Emissão", c["emissao"]))
    if c["prazo"]:
        linhas.append(("Retorno até", c["prazo"]))
    t = d.add_table(rows=len(linhas), cols=2)
    t.autofit = False
    for i, (rot, val) in enumerate(linhas):
        a, b = t.rows[i].cells
        a.width, b.width = Cm(3.6), Cm(13.8)
        for cel in (a, b):
            sombrear(cel._tc.get_or_add_tcPr(), OURO_SUAVE)
            cel.paragraphs[0].paragraph_format.space_after = Pt(1)
            cel.paragraphs[0].paragraph_format.space_before = Pt(1)
        run(a.paragraphs[0], rot.upper(), tam=8.5, negrito=True, cor=CINZA)
        run(b.paragraphs[0], val, tam=10, cor=PRETO)
    tblPr = t._tbl.tblPr
    bordas = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{lado}")
        e.set(qn("w:val"), "nil")
        bordas.append(e)
    esq = OxmlElement("w:left")
    esq.set(qn("w:val"), "single")
    esq.set(qn("w:sz"), "24")
    esq.set(qn("w:color"), hexa(OURO))
    bordas.remove(bordas.find(qn("w:left")))
    bordas.insert(1, esq)
    tblPr.insert_element_before(bordas, "w:shd", "w:tblLayout", "w:tblCellMar",
                                "w:tblLook", "w:tblCaption", "w:tblDescription",
                                "w:tblPrChange")
    d.add_paragraph().paragraph_format.space_after = Pt(4)

    def subtitulo(texto):
        p = d.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run(p, texto.upper(), tam=11.5, negrito=True, cor=OURO, fonte=TITULO, espaco=10)
        borda(p, "bottom", CINZA_LINHA, tam=4, espaco=2)

    def corrido(texto, recuo=0.0):
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if recuo:
            p.paragraph_format.left_indent = Cm(recuo)
        run(p, texto)
        return p

    p = d.add_paragraph()
    run(p, "Prezados(as),")
    for par in c["introducao"]:
        corrido(par)

    subtitulo("Resumo")
    total = len(c["itens"])
    partes = [f"{total} apontamento{'s' if total != 1 else ''}"]
    for pr in PRIORIDADES:
        if c["por_prioridade"].get(pr):
            partes.append(f"{c['por_prioridade'][pr]} de prioridade {pr.lower()}")
    p = d.add_paragraph()
    run(p, " · ".join(partes), tam=10)
    for cat, n in c["por_categoria"].items():
        p = d.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.4)
        p.paragraph_format.space_after = Pt(1)
        run(p, "▪  ", cor=OURO, tam=10)
        run(p, f"{cat}: ", tam=10)
        run(p, str(n), tam=10, cor=CINZA)

    subtitulo("Apontamentos")
    for it in c["itens"]:
        p = d.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.keep_with_next = True
        run(p, f"{it['numero']:02d}   ", tam=13, negrito=True, cor=OURO, fonte=TITULO)
        run(p, it["titulo"], tam=11.5, negrito=True, cor=PRETO)
        p = d.add_paragraph()
        p.paragraph_format.left_indent = Cm(1.05)
        p.paragraph_format.keep_with_next = True
        run(p, it["categoria"] + "   ", tam=8.5, cor=CINZA)
        run(p, f"●  Prioridade {it['prioridade'].lower()}", tam=8.5, negrito=True,
            cor=COR_PRIORIDADE.get(it["prioridade"], OURO))
        for par in it["paragrafos"]:
            corrido(par, recuo=1.05)
        if it["referencia"]:
            p = d.add_paragraph()
            p.paragraph_format.left_indent = Cm(1.05)
            p.paragraph_format.space_after = Pt(1)
            run(p, "Referência: ", tam=9.5, negrito=True, cor=CINZA)
            run(p, it["referencia"], tam=9.5)
        if it["valor"]:
            p = d.add_paragraph()
            p.paragraph_format.left_indent = Cm(1.05)
            p.paragraph_format.space_after = Pt(1)
            run(p, "Valor: ", tam=9.5, negrito=True, cor=CINZA)
            run(p, it["valor"], tam=9.5)
        if it["providencia"]:
            for rotulo, conteudo in (("PROVIDÊNCIA SOLICITADA", None), (None, it["providencia"])):
                p = d.add_paragraph()
                p.paragraph_format.left_indent = Cm(1.25)
                p.paragraph_format.space_after = Pt(0 if rotulo else 6)
                p.paragraph_format.space_before = Pt(4 if rotulo else 0)
                if rotulo:
                    p.paragraph_format.keep_with_next = True
                    run(p, rotulo, tam=8.5, negrito=True, cor=OURO)
                else:
                    run(p, conteudo, tam=10)
                borda(p, "left", OURO, tam=18, espaco=8)
                sombrear(p._p.get_or_add_pPr(), OURO_SUAVE)

    subtitulo("Considerações finais")
    for par in c["encerramento"]:
        corrido(par)
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    run(p, "Atenciosamente,")
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(30)
    p.paragraph_format.space_after = Pt(0)
    run(p, "_" * 42, cor=PRETO)
    if c["responsavel"]:
        p = d.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        run(p, c["responsavel"], negrito=True)
    cargo = c["cargo"] + (f" · CRC {c['crc']}" if c["crc"] else "")
    if cargo.strip():
        p = d.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        run(p, cargo, tam=9.5, cor=CINZA)
    p = d.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    run(p, c["escritorio"].upper(), tam=10, negrito=True, cor=OURO, fonte=TITULO)
    extra = " · ".join(x for x in (f"CNPJ {c['escritorio_cnpj']}" if c["escritorio_cnpj"] else "",
                                   c["escritorio_endereco"]) if x)
    if extra:
        p = d.add_paragraph()
        run(p, extra, tam=8.5, cor=CINZA)
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    run(p, c["data_extenso"], tam=9.5)

    d.core_properties.title = f"{c['titulo']} - {c['empresa']} - {c['competencia']}"
    d.core_properties.author = c["escritorio"]
    destino.parent.mkdir(parents=True, exist_ok=True)
    tmp = destino.with_name(destino.stem + ".tmp.docx")
    d.save(str(tmp))
    _trocar(tmp, destino)


# O Word recusa o arquivo quando os elementos XML fogem da ordem do padrão
# OOXML; por isso cada elemento é inserido antes dos que vêm depois dele.
_SEQ_PPR = ("w:pStyle", "w:keepNext", "w:keepLines", "w:pageBreakBefore", "w:framePr",
            "w:widowControl", "w:numPr", "w:suppressLineNumbers", "w:pBdr", "w:shd",
            "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
            "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN",
            "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
            "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
            "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl",
            "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange")
_DEPOIS_PBDR = _SEQ_PPR[_SEQ_PPR.index("w:pBdr") + 1:]
_DEPOIS_SHD_P = _SEQ_PPR[_SEQ_PPR.index("w:shd") + 1:]
_DEPOIS_SHD_TC = ("w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign",
                  "w:hideMark", "w:headers", "w:cellIns", "w:cellDel", "w:cellMerge",
                  "w:tcPrChange")
_DEPOIS_SPACING_R = ("w:w", "w:kern", "w:position", "w:sz", "w:szCs", "w:highlight",
                     "w:u", "w:effect", "w:bdr", "w:shd", "w:fitText", "w:vertAlign",
                     "w:rtl", "w:cs", "w:em", "w:lang", "w:eastAsianLayout",
                     "w:specVanish", "w:oMath")


def _trocar(tmp: Path, destino: Path) -> None:
    import os
    try:
        os.replace(tmp, destino)
    except PermissionError:
        try:
            tmp.unlink()
        except OSError:
            pass
        raise
