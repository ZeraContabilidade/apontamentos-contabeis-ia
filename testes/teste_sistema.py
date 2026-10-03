# -*- coding: utf-8 -*-
"""Testes do sistema inteiro, sem gastar nada da API.

    python testes/teste_sistema.py

A IA é substituída por um redator falso. O formato do pedido enviado à API
é conferido com a biblioteca oficial apontada para um servidor falso local.
Tudo roda numa pasta temporária: não encosta nos dados reais.
Termina com "RESULTADO: OK".
"""
from __future__ import annotations

import json
import os
import sys
import tempfile
import threading
import time
import traceback
import urllib.error
import urllib.request
import zipfile
from http.server import BaseHTTPRequestHandler, HTTPServer
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ))
TEMP = Path(tempfile.mkdtemp(prefix="apont_teste_"))
os.environ["APONTAMENTOS_DADOS"] = str(TEMP / "dados")

from apontamentos import config, conferencia, documento  # noqa: E402
from apontamentos.base import Base, ErroBase  # noqa: E402
from apontamentos.ia import ErroIA, Redator  # noqa: E402
from apontamentos.servidor import Aplicacao, criar_servidor  # noqa: E402

FALHAS: list[str] = []


def checar(condicao, descricao):
    if condicao:
        print("  ok  ", descricao)
    else:
        print("  FALHOU", descricao)
        FALHAS.append(descricao)


# ----------------------------------------------------------------------
class RedatorFalso:
    """Faz o papel da IA: devolve um texto formal previsível."""
    chamadas = 0
    atraso = 0.0

    def __init__(self, *a, **k):
        pass

    def formalizar(self, original, **k):
        RedatorFalso.chamadas += 1
        time.sleep(RedatorFalso.atraso)
        if "FALHAR" in original:
            raise ErroIA("Sem conexão com a internet (ou a API está fora do ar).")
        return {"titulo": "Apontamento formalizado",
                "texto": "Identificamos o seguinte ponto: " + original.strip() + ".",
                "providencia": "Solicitamos a verificação do item.",
                "pontos_a_confirmar": ["Data do documento não informada"],
                "avisos": conferencia.conferir(original, ["Identificamos: " + original])}

    def testar(self):
        return "ok"


class _Bloco:
    def __init__(self, tipo, texto=""):
        self.type, self.text = tipo, texto


class _Resposta:
    def __init__(self, texto, parada="end_turn"):
        self.content = [_Bloco("text", texto)]
        self.stop_reason = parada


class ClienteFalso:
    """Imita anthropic.Anthropic para testar o Redator sem rede."""

    def __init__(self, resposta, erro_reserva=None):
        self.pedidos = []
        self._resposta = resposta
        self._erro_reserva = erro_reserva
        externo = self

        class _Msgs:
            def create(self_inner, **k):
                externo.pedidos.append(("normal", k))
                return externo._resposta

        class _BetaMsgs:
            def create(self_inner, **k):
                externo.pedidos.append(("beta", k))
                if externo._erro_reserva:
                    raise externo._erro_reserva
                return externo._resposta

        class _Beta:
            messages = _BetaMsgs()

        self.messages = _Msgs()
        self.beta = _Beta()


class Erro400(Exception):
    status_code = 400


# ----------------------------------------------------------------------
def teste_conferencia():
    print("\n[conferência de números]")
    c = conferencia.conferir
    checar(c("nf 1234 de 10/09 valor 1500", ["Nota Fiscal nº 1234, de 10/09/2026, valor de R$ 1.500,00"],
             ["2026-09"]) == [], "mesmo valor em formatos diferentes passa")
    av = c("nf 1234 valor 1500", ["Nota 1234 de R$ 1.600,00"])
    checar(any("1.600,00" in a for a in av), "valor inventado é apontado")
    checar(any("1500" in a for a in av), "valor do original que sumiu é apontado")
    checar(c("cnpj 12.345.678/0001-90", ["CNPJ 12.345.678/0001-90"]) == [], "CNPJ preservado passa")
    checar(any("10" in a for a in c("pagar", ["multa de 10%"])), "percentual inventado é apontado")
    checar(c("competência 09/2026", ["competência de setembro de 2026"]) != [] or True,
           "mês por extenso não quebra a conferência")
    checar(c("venceu 05/09/26", ["vencido em 05/09/2026"]) == [], "ano com 2 dígitos = 4 dígitos")


def teste_redator():
    print("\n[redator (IA) com cliente falso]")
    ok = json.dumps({"titulo": "Nota fiscal sem comprovante.", "texto": "Identificamos a **Nota Fiscal** nº 1234.",
                     "providencia": "Solicitamos o envio do comprovante.", "pontos_a_confirmar": ["Valor não informado"]})
    cli = ClienteFalso(_Resposta(ok))
    r = Redator("", "claude-opus-5-5", "medium", cliente=cli).formalizar("nf 1234 sem comprovante")
    checar(r["titulo"] == "Nota fiscal sem comprovante", "título sem ponto final")
    checar("**" not in r["texto"], "markdown removido")
    checar(r["avisos"] == [], "sem aviso quando os números batem")
    tipo, pedido = cli.pedidos[0]
    checar(tipo == "beta" and pedido.get("fallbacks") == "default", "Opus 5.5 usa a reserva automática")
    checar(pedido["output_config"]["effort"] == "medium", "esforço enviado")
    checar(pedido["output_config"]["format"]["type"] == "json_schema", "resposta estruturada pedida")
    checar("Anotação do contador" in pedido["messages"][0]["content"], "texto do contador enviado")

    cli = ClienteFalso(_Resposta(ok))
    Redator("", "claude-haiku-4-5", "medium", cliente=cli).formalizar("x")
    tipo, pedido = cli.pedidos[0]
    checar(tipo == "normal" and "effort" not in pedido["output_config"],
           "Haiku sem esforço e sem reserva (não aceita)")

    cli = ClienteFalso(_Resposta(ok), erro_reserva=Erro400("fallbacks not available"))
    red = Redator("", "claude-opus-5-5", "medium", cliente=cli)
    red.formalizar("x")
    red.formalizar("y")
    checar([t for t, _ in cli.pedidos] == ["beta", "normal", "normal"],
           "conta sem reserva: segue sem ela e não insiste")

    for resposta, trecho in ((_Resposta("{}", "refusal"), "recusou"),
                             (_Resposta("não é json"), "formato"),
                             (_Resposta(json.dumps({"titulo": "", "texto": "", "providencia": "",
                                                    "pontos_a_confirmar": []})), "vazio"),
                             (_Resposta("{", "max_tokens"), "cortada")):
        try:
            Redator("", "claude-opus-5-5", cliente=ClienteFalso(resposta)).formalizar("x")
            checar(False, f"erro '{trecho}' detectado")
        except ErroIA as e:
            checar(trecho in str(e), f"erro '{trecho}' detectado")

    try:
        Redator("", "claude-opus-5-5")
        checar(False, "sem chave avisa")
    except ErroIA as e:
        checar("chave" in str(e).lower(), "sem chave avisa")


class _ApiFalsa(BaseHTTPRequestHandler):
    recebidos: list = []

    def log_message(self, *a):
        pass

    def do_POST(self):
        corpo = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
        _ApiFalsa.recebidos.append((self.path, dict(self.headers), corpo))
        resposta = {
            "id": "msg_teste", "type": "message", "role": "assistant", "model": corpo["model"],
            "content": [{"type": "text", "text": json.dumps({
                "titulo": "Nota fiscal sem comprovante de pagamento",
                "texto": "Identificamos a Nota Fiscal nº 1234 sem o comprovante de pagamento.",
                "providencia": "Solicitamos o envio do comprovante.", "pontos_a_confirmar": []})}],
            "stop_reason": "end_turn", "stop_sequence": None,
            "usage": {"input_tokens": 10, "output_tokens": 10}}
        dados = json.dumps(resposta).encode()
        self.send_response(200)
        self.send_header("Content-Type", "application/json")
        self.send_header("Content-Length", str(len(dados)))
        self.end_headers()
        self.wfile.write(dados)


def teste_pedido_real_da_biblioteca():
    print("\n[pedido montado pela biblioteca oficial da Anthropic]")
    import anthropic
    srv = HTTPServer(("127.0.0.1", 0), _ApiFalsa)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    try:
        red = Redator("sk-teste", "claude-opus-5-5", "medium", "wrkspc_teste123")
        red.cliente = red.cliente.with_options(base_url=f"http://127.0.0.1:{srv.server_port}",
                                               max_retries=0)
        r = red.formalizar("nf 1234 sem comprovante", categoria="Documentação pendente")
        caminho, cab, corpo = _ApiFalsa.recebidos[-1]
        checar(caminho.startswith("/v1/messages"), "chama /v1/messages")
        checar("server-side-fallback-2026-07-01" in cab.get("anthropic-beta", ""), "cabeçalho beta da reserva")
        checar(cab.get("anthropic-workspace-id") == "wrkspc_teste123", "Workspace ID enviado no cabeçalho")
        red2 = Redator("sk-teste", "claude-opus-5-5", "medium", "")
        red2.cliente = red2.cliente.with_options(base_url=f"http://127.0.0.1:{srv.server_port}",
                                                 max_retries=0)
        red2.formalizar("x")
        checar("anthropic-workspace-id" not in {k.lower() for k in _ApiFalsa.recebidos[-1][1]},
               "sem Workspace ID: cabeçalho não vai")
        erro = anthropic.BadRequestError(
            "This API key is not scoped to a workspace, so this request must include the "
            "anthropic-workspace-id header", response=__import__("httpx2").Response(
                400, request=__import__("httpx2").Request("POST", "http://x")), body=None)
        checar("Workspace ID" in red._mensagem_erro(erro), "erro de workspace explicado")
        checar(corpo.get("fallbacks") == "default", "fallbacks=default no corpo")
        checar(corpo["model"] == "claude-opus-5-5", "modelo Opus 5.5")
        checar(corpo["output_config"]["format"]["schema"]["additionalProperties"] is False,
               "esquema JSON estrito")
        checar("thinking" not in corpo, "não desliga o raciocínio (proibido no Opus 5.5)")
        checar(r["titulo"].startswith("Nota fiscal"), "resposta lida")
    finally:
        srv.shutdown()


# ----------------------------------------------------------------------
_SEQ_RPR = ("w:rStyle", "w:rFonts", "w:b", "w:bCs", "w:i", "w:iCs", "w:caps", "w:smallCaps",
            "w:strike", "w:dstrike", "w:outline", "w:shadow", "w:emboss", "w:imprint",
            "w:noProof", "w:snapToGrid", "w:vanish", "w:webHidden", "w:color", "w:spacing",
            "w:w", "w:kern", "w:position", "w:sz", "w:szCs", "w:highlight", "w:u", "w:effect",
            "w:bdr", "w:shd", "w:fitText", "w:vertAlign", "w:rtl", "w:cs", "w:em", "w:lang",
            "w:eastAsianLayout", "w:specVanish", "w:oMath")
_SEQ_TBLPR = ("w:tblStyle", "w:tblpPr", "w:tblOverlap", "w:bidiVisual", "w:tblStyleRowBandSize",
              "w:tblStyleColBandSize", "w:tblW", "w:jc", "w:tblCellSpacing", "w:tblInd",
              "w:tblBorders", "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook")
_SEQ_TCPR = ("w:cnfStyle", "w:tcW", "w:gridSpan", "w:hMerge", "w:vMerge", "w:tcBorders",
             "w:shd", "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark")


def _ordem_xml_ok(caminho: Path) -> list[str]:
    """O Word recusa o arquivo se os elementos fogem da ordem do padrão."""
    from lxml import etree
    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    seqs = {"pPr": documento._SEQ_PPR, "rPr": _SEQ_RPR, "tblPr": _SEQ_TBLPR, "tcPr": _SEQ_TCPR}
    problemas = []
    with zipfile.ZipFile(caminho) as z:
        for nome in z.namelist():
            if not (nome.startswith("word/") and nome.endswith(".xml")) or "styles" in nome:
                continue
            raiz = etree.fromstring(z.read(nome))
            for pai, seq in seqs.items():
                for e in raiz.iter(W + pai):
                    if pai == "rPr" and e.getparent() is not None and e.getparent().tag == W + "pPr":
                        continue
                    nomes = ["w:" + f.tag.replace(W, "") for f in e]
                    pos = [seq.index(n) for n in nomes if n in seq]
                    if pos != sorted(pos):
                        problemas.append(f"{nome}: {pai} fora de ordem {nomes}")
    return problemas


def _cfg_teste() -> config.Config:
    cfg = config.Config()
    cfg.chave_api = "sk-teste"
    cfg.pasta_saida = str(TEMP / "saida")
    cfg.responsavel_nome = "Responsável Teste"
    cfg.responsavel_crc = "1SP000000/O-0"
    cfg.escritorio_telefone = "(11) 0000-0000"
    return cfg


def teste_aplicacao():
    print("\n[fluxo completo: base, IA falsa, edição, documento]")
    app = Aplicacao(base=Base(TEMP / "a.db"), cfg=_cfg_teste(),
                    fabrica_redator=lambda c: RedatorFalso(), salvar_config=False)
    try:
        app.criar_documento({"empresa": "", "competencia": "09/2026"})
        checar(False, "empresa obrigatória")
    except ErroBase:
        checar(True, "empresa obrigatória")
    try:
        app.criar_documento({"empresa": "X", "competencia": "13/2026"})
        checar(False, "competência inválida recusada")
    except ErroBase:
        checar(True, "competência inválida recusada")

    doc = app.criar_documento({"empresa": "Mercado Bom Preço Ltda", "cnpj": "12.345.678/0001-90",
                               "competencia": "09/2026", "prazo_retorno": "2026-10-15"})
    checar(doc["competencia"] == "2026-09" and doc["competencia_extenso"] == "Setembro/2026",
           "competência normalizada")
    ids = []
    for texto, cat, prio in (("nf 1234 do fornecedor Alfa sem boleto, valor 1500", "Documentação pendente", "Alta"),
                             ("extrato com diferença de 230,15 em 30/09", "Conciliação bancária", "Média"),
                             ("FALHAR teste de queda da IA", "Outros", "Baixa")):
        r = app.adicionar(doc["id"], {"original": texto, "categoria": cat, "prioridade": prio})
        checar(r["apontamento"]["situacao"] == "pendente", "novo apontamento entra como pendente")
        ids.append(r["apontamento"]["id"])
    try:
        documento.gerar(app.base.obter_documento(doc["id"]), app.cfg)
        checar(False, "não gera documento com item pendente")
    except documento.ErroDocumento:
        checar(True, "não gera documento com item pendente")

    r = app.formalizar(ids[0])
    ap = r["apontamento"]
    checar(ap["situacao"] == "formalizado" and ap["titulo"], "formalizado pela IA")
    checar(ap["original"].startswith("nf 1234"), "texto original guardado intacto")
    checar(any(a.startswith("Sugestão da IA") for a in ap["avisos"]), "sugestões da IA viram lembrete")
    app.formalizar(ids[1])
    r = app.formalizar(ids[2])
    checar(r["apontamento"]["situacao"] == "erro" and r["erro"], "queda da IA: item fica com o texto original")
    checar(r["apontamento"]["texto"].startswith("FALHAR"), "queda da IA não perde o texto")

    r = app.editar(ids[1], {"texto": "Verificamos diferença de R$ 230,15 no saldo bancário de 30/09/2026.",
                            "titulo": "Diferença na conciliação bancária"})
    checar(r["apontamento"]["situacao"] == "editado" and r["apontamento"]["avisos"] == [], "edição manual")
    try:
        app.editar(ids[1], {"texto": "  "})
        checar(False, "texto vazio recusado")
    except Exception:
        checar(True, "texto vazio recusado")
    app.mover(ids[2], -1)
    ordem = [a["id"] for a in app.base.obter_documento(doc["id"])["apontamentos"]]
    checar(ordem == [ids[0], ids[2], ids[1]], "mover para cima")
    app.usar_original(ids[2])
    checar(app.base.obter_apontamento(ids[2])["situacao"] == "sem_ia", "usar meu texto")
    app.excluir_apontamento(ids[2])
    d = app.base.obter_documento(doc["id"])
    checar([a["ordem"] for a in d["apontamentos"]] == [1, 2], "renumera depois de excluir")

    sem_chave = Aplicacao(base=app.base, cfg=config.Config(), salvar_config=False,
                          fabrica_redator=lambda c: (_ for _ in ()).throw(AssertionError("chamou IA")))
    r = sem_chave.adicionar(doc["id"], {"original": "sem chave", "categoria": "Outros"})
    r = sem_chave.formalizar(r["apontamento"]["id"])
    checar(r["apontamento"]["situacao"] == "sem_ia", "sem chave: não chama a IA, usa o texto")
    app.excluir_apontamento(r["apontamento"]["id"])

    r = app.finalizar(doc["id"], ["docx", "pdf"])
    arquivos = [Path(a) for a in r["arquivos"]]
    checar(len(arquivos) == 2 and all(a.exists() and a.stat().st_size > 5000 for a in arquivos),
           "DOCX e PDF gravados")
    checar(arquivos[0].parent.name == "Mercado Bom Preço Ltda", "subpasta por empresa")
    checar(r["documento"]["status"] == "finalizado", "documento finalizado")
    try:
        app.adicionar(doc["id"], {"original": "depois de finalizar"})
        checar(False, "finalizado não aceita alteração")
    except ErroBase:
        checar(True, "finalizado não aceita alteração")

    from docx import Document
    dx = Document(str(arquivos[0]))
    tudo = "\n".join(p.text for p in dx.paragraphs)
    checar("Identificamos o seguinte ponto: nf 1234" in tudo, "DOCX traz o texto formal")
    checar("Solicitamos o retorno até 15/10/2026" in tudo, "DOCX traz o prazo de retorno")
    checar("Responsável Teste" in tudo, "DOCX traz a assinatura")
    checar(not _ordem_xml_ok(arquivos[0]), "DOCX com XML na ordem do padrão (abre no Word)")
    for p in _ordem_xml_ok(arquivos[0])[:5]:
        print("        ", p)
    with zipfile.ZipFile(arquivos[0]) as z:
        checar(any(n.startswith("word/media/") for n in z.namelist()), "DOCX com a logo")

    from pypdf import PdfReader
    pdf = PdfReader(str(arquivos[1]))
    texto_pdf = "\n".join(p.extract_text() for p in pdf.pages)
    checar("Mercado Bom Preço Ltda" in texto_pdf and "Setembro/2026" in texto_pdf, "PDF com empresa e competência")
    checar("PROVIDÊNCIA SOLICITADA" in texto_pdf, "PDF com a providência")
    checar("Página 1 de" in texto_pdf, "PDF com numeração de páginas")

    app.reabrir(doc["id"])
    checar(app.base.obter_documento(doc["id"])["status"] == "rascunho", "reabrir para edição")
    lista = {d["id"]: d for d in app.base.listar_documentos()}
    checar(lista[doc["id"]]["quantidade_alta"] == 1, "painel conta os de prioridade alta")

    # trazer pendências para o mês seguinte
    outubro = app.criar_documento({"empresa": "Mercado Bom Preço Ltda", "competencia": "10/2026"})
    r = app.copiar_de(outubro["id"], [ids[0]])
    copiado = app.base.obter_documento(outubro["id"])["apontamentos"]
    checar(r["copiados"] == 1 and copiado[0]["texto"] == app.base.obter_apontamento(ids[0])["texto"],
           "trazer apontamento de outro mês")
    checar(any(a.startswith("Trazido de Setembro/2026") for a in copiado[0]["avisos"]), "aviso de onde veio")
    checar(copiado[0]["original"].startswith("nf 1234"), "trazido mantém o texto original")

    # backup: exporta e importa em outra base (outro aparelho)
    bk = app.base.exportar()
    checar(bk["app"] == "apontamentos-contabeis-ia" and len(bk["documentos"]) == 2, "backup exporta tudo")
    outra = Base(TEMP / "outro_aparelho.db")
    r = outra.importar(json.loads(json.dumps(bk)))
    checar(r == {"importados": 2, "ignorados": 0}, "backup importado em outro aparelho")
    checar(outra.importar(bk) == {"importados": 0, "ignorados": 2}, "importar de novo não duplica")
    trazido = [d for d in outra.listar_documentos() if d["competencia"] == "2026-09"][0]
    checar(outra.obter_documento(trazido["id"])["apontamentos"][0]["original"].startswith("nf 1234"),
           "backup preserva o texto original")

    # documento longo: quebra de página sem erro
    longo = app.criar_documento({"empresa": "Empresa: Teste / Longa?", "competencia": "2026-08"})
    for n in range(25):
        a = app.adicionar(longo["id"], {"original": f"item {n} " + "texto longo " * 40, "prioridade": "Alta"})
        app.formalizar(a["apontamento"]["id"])
    r = app.finalizar(longo["id"], ["pdf", "docx"])
    checar(all(Path(a).exists() for a in r["arquivos"]), "documento com 25 itens e nome com caracteres proibidos")
    checar(len(PdfReader(r["arquivos"][0]).pages) > 3, "PDF longo em várias páginas")


def teste_config():
    print("\n[configurações]")
    cfg = config.carregar()
    cfg.atualizar({"chave_api": "sk-ant-segredo", "modelo": "modelo-inexistente", "esforco": "low",
                   "workspace_id": "  wrkspc_abc  ",
                   "escritorio_nome": "Zera", "campo_estranho": 1})
    config.salvar(cfg)
    bruto = json.loads(config.arquivo_config().read_text(encoding="utf-8"))
    if os.name == "nt":
        checar(bruto["chave_api"].startswith("dpapi:"), "chave gravada protegida")
    lido = config.carregar()
    checar(lido.chave_api == "sk-ant-segredo", "chave lida de volta")
    checar(lido.modelo == config.MODELO_PADRAO, "modelo inválido ignorado")
    checar(lido.esforco == "low", "esforço salvo")
    checar(lido.workspace_id == "wrkspc_abc", "Workspace ID salvo sem espaços")
    checar("chave_api" not in lido.publico() and lido.publico()["chave_configurada"],
           "a chave nunca vai para a página")
    lido.atualizar({"chave_api": ""})
    checar(lido.chave_api == "sk-ant-segredo", "campo vazio mantém a chave")


# ----------------------------------------------------------------------
def _pedir(base, caminho, corpo=None, cabecalho=True, host=None):
    dados = None if corpo is None else json.dumps(corpo).encode()
    req = urllib.request.Request(base + caminho, data=dados, method="POST" if corpo is not None else "GET")
    if corpo is not None:
        req.add_header("Content-Type", "application/json")
        if cabecalho:
            req.add_header("X-Requested-With", "Apontamentos")
    if host:
        req.add_header("Host", host)
    try:
        with urllib.request.urlopen(req, timeout=20) as r:
            return r.status, r.read(), dict(r.headers)
    except urllib.error.HTTPError as e:
        return e.code, e.read(), dict(e.headers)


def teste_servidor():
    print("\n[servidor local]")
    app = Aplicacao(base=Base(TEMP / "s.db"), cfg=_cfg_teste(),
                    fabrica_redator=lambda c: RedatorFalso(), salvar_config=False)
    srv = criar_servidor(app, 0)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    base = f"http://127.0.0.1:{srv.server_port}"
    try:
        st, corpo, cab = _pedir(base, "/")
        checar(st == 200 and b"app.js" in corpo, "página abre")
        checar("script-src 'self'" in cab.get("Content-Security-Policy", ""), "página com CSP")
        for arq in ("/app.js", "/estilo.css", "/marca/logo_claro.png", "/marca/emblema.png",
                    "/local/servidor-local.js", "/vendor/jspdf.umd.min.js", "/sw.js"):
            checar(_pedir(base, arq)[0] == 200, f"serve {arq}")
        st, _, cab = _pedir(base, "/manifest.webmanifest")
        checar(st == 200 and "manifest" in cab.get("Content-Type", ""), "manifesto do aplicativo")
        checar(_pedir(base, "/../apontamentos/servidor.py")[0] == 404, "não sai da pasta docs")
        checar(_pedir(base, "/marca/..%2f..%2fapontamentos%2fservidor.py")[0] == 404, "não sai da pasta (codificado)")
        checar(_pedir(base, "/README.md")[0] == 404, "não serve .md")
        from apontamentos import __version__
        checar(f'VERSAO_APP = "{__version__}"' in (RAIZ / "docs" / "app.js").read_text(encoding="utf-8"),
               "aplicativo e programa na mesma versão")
        st, corpo, _ = _pedir(base, "/api/ping")
        checar(st == 200 and json.loads(corpo)["app"] == "apontamentos", "ping identifica o programa")
        st, corpo, _ = _pedir(base, "/api/estado")
        estado = json.loads(corpo)
        checar(st == 200 and "chave_api" not in json.dumps(estado["config"]), "estado sem a chave")
        checar(_pedir(base, "/api/documentos", {"empresa": "X", "competencia": "09/2026"},
                      cabecalho=False)[0] == 403, "POST sem cabeçalho recusado")
        checar(_pedir(base, "/api/estado", host="exemplo.com")[0] == 403, "Host de fora recusado")
        st, corpo, _ = _pedir(base, "/api/documentos", {"empresa": "Padaria Teste", "competencia": "09/2026"})
        doc = json.loads(corpo)
        checar(st == 200 and doc["id"], "cria documento")
        st, corpo, _ = _pedir(base, f"/api/documentos/{doc['id']}/apontamentos",
                              {"original": "nf 99 sem xml", "categoria": "Fiscal / Tributos"})
        ap = json.loads(corpo)["apontamento"]
        st, corpo, _ = _pedir(base, f"/api/apontamentos/{ap['id']}/formalizar", {})
        checar(json.loads(corpo)["apontamento"]["situacao"] == "formalizado", "formaliza via HTTP")
        checar(_pedir(base, "/api/documentos/999999")[0] == 400, "documento inexistente")
        checar(_pedir(base, "/api/qualquer", {})[0] == 404, "rota desconhecida")
        st, corpo, _ = _pedir(base, f"/api/documentos/{doc['id']}/finalizar", {"formatos": ["pdf"]})
        checar(st == 200, "finaliza via HTTP")
        st, corpo, cab = _pedir(base, f"/api/arquivo?doc={doc['id']}&fmt=pdf")
        checar(st == 200 and corpo.startswith(b"%PDF") and "attachment" in cab.get("Content-Disposition", ""),
               "baixa o PDF")
        checar(_pedir(base, f"/api/arquivo?doc={doc['id']}&fmt=docx")[0] == 404, "DOCX não gerado: 404")
        st, corpo, _ = _pedir(base, "/api/backup")
        bk = json.loads(corpo)
        checar(st == 200 and bk["app"] == "apontamentos-contabeis-ia" and "sk-teste" not in corpo.decode(),
               "backup via HTTP, sem a chave")
        st, corpo, _ = _pedir(base, "/api/backup/importar", bk)
        checar(st == 200 and json.loads(corpo)["ignorados"] == 1, "importar o mesmo backup não duplica")
        checar(_pedir(base, "/api/backup/importar", {"app": "outro"})[0] == 400, "backup de outro sistema recusado")
        st, corpo, _ = _pedir(base, "/api/documentos", {"empresa": "Padaria Teste", "competencia": "10/2026"})
        novo = json.loads(corpo)
        st, corpo, _ = _pedir(base, f"/api/documentos/{novo['id']}/copiar_de", {"ids": [ap["id"]]})
        checar(st == 200 and json.loads(corpo)["copiados"] == 1, "trazer apontamento de outro mês via HTTP")
    finally:
        srv.shutdown()
        srv.server_close()


def teste_tela():
    """Usa o Chromium para operar a tela de verdade (pula se não houver)."""
    print("\n[tela no navegador]")
    try:
        from playwright.sync_api import sync_playwright
    except ImportError:
        print("  (pulado: playwright não instalado)")
        return
    app = Aplicacao(base=Base(TEMP / "t.db"), cfg=_cfg_teste(),
                    fabrica_redator=lambda c: RedatorFalso(), salvar_config=False)
    RedatorFalso.atraso = 0.8
    srv = criar_servidor(app, 0)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    base = f"http://127.0.0.1:{srv.server_port}/"
    erros = []
    try:
        with sync_playwright() as p:
            caminho = os.environ.get("CHROMIUM") or ("/opt/pw-browsers/chromium-1194/chrome-linux/chrome"
                                                    if Path("/opt/pw-browsers/chromium-1194/chrome-linux/chrome").exists() else None)
            nav = p.chromium.launch(executable_path=caminho) if caminho else p.chromium.launch()
            pg = nav.new_page(viewport={"width": 1500, "height": 950}, bypass_csp=True)
            pg.on("pageerror", lambda e: erros.append(str(e)))
            pg.on("console", lambda m: erros.append(m.text) if m.type == "error" else None)
            pg.on("dialog", lambda d: d.accept())
            pg.goto(base)
            pg.wait_for_selector(".painel")
            checar(pg.locator(".tile").count() == 4, "painel com indicadores")
            pg.click("#bt-novo")
            pg.fill("input[list=empresas-conhecidas]", "Oficina Teste ME")
            pg.fill("input[type=month]", "2026-09")
            pg.click("text=Criar documento")
            pg.wait_for_selector("#novo-texto")
            pg.fill("#novo-texto", "nf 555 lancada duas vezes, valor 980,00")
            pg.select_option("#novo-prioridade", "Alta")
            pg.keyboard.press("Control+Enter")
            pg.wait_for_selector(".ap.pendente", timeout=5000)
            checar(pg.input_value("#novo-texto") == "", "caixa limpa para o próximo apontamento")
            pg.fill("#novo-texto", "segundo apontamento enquanto a IA trabalha")
            pg.click("text=Adicionar apontamento")
            pg.wait_for_function("document.querySelectorAll('.ap.formalizado').length === 2", timeout=10000)
            checar(pg.locator(".papel .p-item").count() == 2, "prévia mostra os 2 apontamentos")
            checar("Identificamos o seguinte ponto: nf 555" in pg.inner_text(".papel"), "prévia com o texto formal")
            pg.locator(".ap").first.locator("text=Editar").click()
            pg.locator(".ap.edicao textarea").first.fill("Texto ajustado à mão pelo contador.")
            pg.locator(".ap.edicao").locator("text=Salvar").first.click()
            pg.wait_for_selector(".ap.editado")
            checar("Texto ajustado à mão" in pg.inner_text(".papel"), "edição aparece na prévia")
            pg.locator(".ap").nth(1).locator("text=↑").click()
            pg.wait_for_function("document.querySelector('.ap .ap-titulo').textContent !== ''")
            pg.click("text=Finalizar e gerar")
            pg.wait_for_selector(".faixa-final", timeout=15000)
            checar(pg.locator("text=Baixar PDF").count() == 1, "finalizado com botões de baixar")
            with pg.expect_download() as baixa:
                pg.click("text=Baixar PDF")
            checar(baixa.value.suggested_filename.endswith(".pdf"), "download do PDF")
            pg.click("text=Mensagem para o cliente")
            msg = pg.input_value(".modal textarea")
            checar("Oficina Teste ME" in msg and "Setembro/2026" in msg, "mensagem pronta para o cliente")
            pg.click(".modal >> text=Fechar")
            pg.screenshot(path=str(TEMP / "tela.png"), full_page=False)
            pg.click("#bt-config")
            pg.wait_for_selector("text=Salvar configurações")
            nav.close()
        checar(not erros, "sem erro de JavaScript na tela")
        for e in erros[:5]:
            print("        ", e)
        print("   (captura da tela em", TEMP / "tela.png", ")")
    finally:
        RedatorFalso.atraso = 0.0
        srv.shutdown()
        srv.server_close()


# ----------------------------------------------------------------------
_NODE_PARIDADE = r"""
const fs = require("fs");
globalThis.window = globalThis;
for (const f of ["conteudo.js", "conferencia.js", "ia.js"])
  eval(fs.readFileSync(process.argv[1] + "/" + f, "utf8"));
const entrada = JSON.parse(fs.readFileSync(0, "utf8"));
const saida = {
  instrucoes: IA.INSTRUCOES, esquema: IA.ESQUEMA,
  conferencias: entrada.casos.map(c => Conferencia.conferir(c[0], c[1], c[2])),
  conteudo: Conteudo.montarConteudo(entrada.doc, entrada.cfg, new Date(2026, 9, 2)),
  valores: entrada.valores.map(v => Conteudo.formatarValor(v)),
};
process.stdout.write(JSON.stringify(saida));
"""


def teste_paridade_js():
    """O aplicativo (iPhone) usa JavaScript; confere que faz o mesmo que o Python."""
    print("\n[aplicativo x programa: mesmas regras]")
    import shutil
    import subprocess
    from datetime import date
    from apontamentos import ia
    if not shutil.which("node"):
        print("  (pulado: node não instalado)")
        return
    casos = [
        ["nf 1234 lançada em 10/09, valor 1500", ["Nota fiscal 1234 lançada em 10/09/2026, no valor de R$ 1.500,00."], ["Setembro/2026"]],
        ["diferença de 230,15", ["Diferença de R$ 230,15 e multa de 2%."], []],
        ["nf 55 e nf 66", ["Nota fiscal 55."], []],
        ["pagamento 05/2026 de 1.234,56", ["Pagamento de maio (05/2026) de R$ 1.234,56 em 2026-05-10."], []],
        ["valor 12.5", ["Valor de 12,50."], ["99"]],
    ]
    cfg = _cfg_teste()
    cfg.escritorio_email = "contato@exemplo.com"
    doc = {"empresa": "Loja Teste", "cnpj": "12.345.678/0001-90", "competencia": "2026-09",
           "destinatario": "Sr. José", "prazo_retorno": "2026-10-15", "apontamentos": [
               {"id": 1, "situacao": "formalizado", "titulo": "Nota em duplicidade", "categoria": "Fiscal / Tributos",
                "prioridade": "Alta", "texto": "Primeiro parágrafo.\n\nSegundo\nparágrafo.", "original": "x",
                "providencia": " Solicitamos o estorno. ", "valor": "1500", "referencia": "NF 1234"},
               {"id": 2, "situacao": "sem_ia", "titulo": "", "categoria": "Outros", "prioridade": "Baixa",
                "texto": "", "original": "texto do contador", "providencia": "", "valor": "", "referencia": ""}]}
    valores = ["1500", "1.500,00", "R$ 980,5", "12.5", "abc", "", "1234567,891"]
    entrada = {"casos": casos, "doc": doc, "valores": valores,
               "cfg": {k: getattr(cfg, k) for k in vars(cfg)}}
    r = subprocess.run(["node", "-e", _NODE_PARIDADE, str(RAIZ / "docs" / "local")],
                       input=json.dumps(entrada), capture_output=True, text=True, timeout=60)
    if r.returncode:
        print(r.stderr[-2000:])
        checar(False, "JavaScript do aplicativo carrega")
        return
    js = json.loads(r.stdout)
    checar(js["instrucoes"] == ia.INSTRUCOES, "mesmas instruções para a IA")
    checar(js["esquema"] == ia.ESQUEMA, "mesmo formato de resposta da IA")
    for n, (c, saida) in enumerate(zip(casos, js["conferencias"]), 1):
        py = conferencia.conferir(c[0], c[1], c[2])
        checar(py == saida, f"mesma conferência de números (caso {n})")
        if py != saida:
            print("         python:", py, "\n         js:    ", saida)
    py = documento.montar_conteudo(doc, cfg, date(2026, 10, 2))
    for k in py:
        if py[k] != js["conteudo"].get(k):
            print("         diferença em", k, ":", py[k], "x", js["conteudo"].get(k))
    checar(all(py[k] == js["conteudo"].get(k) for k in py), "mesmo conteúdo do documento")
    checar([documento.formatar_valor(v) for v in valores] == js["valores"], "mesma formatação de valores")


class _ServidorEstatico(BaseHTTPRequestHandler):
    """Faz o papel do GitHub Pages: só arquivos, nada de /api/."""
    def log_message(self, *a):
        pass

    def do_GET(self):
        from urllib.parse import unquote, urlparse
        caminho = unquote(urlparse(self.path).path).lstrip("/") or "index.html"
        alvo = (RAIZ / "docs" / caminho).resolve()
        if not alvo.is_file() or (RAIZ / "docs") not in alvo.parents:
            self.send_response(404)
            self.send_header("Content-Type", "text/html")
            self.end_headers()
            self.wfile.write(b"<h1>404</h1>")
            return
        tipos = {".js": "text/javascript", ".css": "text/css", ".html": "text/html; charset=utf-8",
                 ".png": "image/png", ".ttf": "font/ttf", ".webmanifest": "application/manifest+json"}
        dados = alvo.read_bytes()
        self.send_response(200)
        self.send_header("Content-Type", tipos.get(alvo.suffix, "application/octet-stream"))
        self.send_header("Content-Length", str(len(dados)))
        self.end_headers()
        self.wfile.write(dados)


def teste_aplicativo():
    """O aplicativo como no iPhone: sem o programa, dados no aparelho, IA direto na API."""
    print("\n[aplicativo (iPhone/iPad): tela de celular, sem o programa]")
    try:
        from playwright.sync_api import sync_playwright
    except ImportError:
        print("  (pulado: playwright não instalado)")
        return
    from http.server import ThreadingHTTPServer
    srv = ThreadingHTTPServer(("127.0.0.1", 0), _ServidorEstatico)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    base = f"http://127.0.0.1:{srv.server_port}/"
    pedidos_ia = []

    def api_falsa(route):
        req = route.request
        corpo = json.loads(req.post_data or "{}")
        pedidos_ia.append({"cab": req.headers, "corpo": corpo})
        anotacao = corpo["messages"][0]["content"].split("Anotação do contador:\n", 1)[-1]
        resposta = {"titulo": "Apontamento formalizado", "texto": "Identificamos o seguinte ponto: " + anotacao + ".",
                    "providencia": "Solicitamos a verificação do item.", "pontos_a_confirmar": []}
        if "INVENTAR" in anotacao:
            resposta["texto"] += " Multa de 20%."
        time.sleep(0.4)
        route.fulfill(status=200, content_type="application/json", headers={"access-control-allow-origin": "*"},
                      body=json.dumps({"content": [{"type": "text", "text": json.dumps(resposta)}],
                                       "stop_reason": "end_turn"}))

    erros = []
    try:
        with sync_playwright() as p:
            caminho = os.environ.get("CHROMIUM") or ("/opt/pw-browsers/chromium-1194/chrome-linux/chrome"
                                                    if Path("/opt/pw-browsers/chromium-1194/chrome-linux/chrome").exists() else None)
            nav = p.chromium.launch(executable_path=caminho) if caminho else p.chromium.launch()
            ctx = nav.new_context(viewport={"width": 390, "height": 844}, is_mobile=True, has_touch=True,
                                  device_scale_factor=2, accept_downloads=True, service_workers="block",
                                  user_agent="Mozilla/5.0 (iPhone; CPU iPhone OS 18_0 like Mac OS X) "
                                             "AppleWebKit/605.1.15 (KHTML, like Gecko) Version/18.0 Mobile/15E148 Safari/604.1")
            ctx.route("https://api.anthropic.com/**", api_falsa)
            pg = ctx.new_page()
            pg.on("pageerror", lambda e: erros.append(str(e)))
            pg.on("console", lambda m: erros.append(m.text) if m.type == "error" and "404" not in m.text else None)
            pg.on("dialog", lambda d: d.accept())
            pg.goto(base)
            pg.wait_for_selector(".painel")
            checar("aplicativo" in pg.inner_text("#versao"), "abre em modo aplicativo (sem o programa)")
            checar(pg.locator("text=Adicionar à Tela de Início").count() >= 1, "ensina a instalar no iPhone")
            largura = pg.evaluate("document.documentElement.scrollWidth")
            checar(largura <= 392, f"cabe na tela do celular sem rolar para o lado ({largura}px)")

            # configurações: chave guardada no aparelho
            pg.click("#bt-config")
            pg.fill("input[type=password]", "sk-ant-teste-aparelho")
            pg.locator("input.campo").nth(4).fill("Zera Contabilidade")
            pg.click("text=Salvar configurações")
            pg.wait_for_selector("text=Chave cadastrada neste aparelho")
            checar("sk-ant-teste" not in pg.content(), "a chave não volta para a tela")

            # menu lateral vira gaveta
            pg.click("#bt-menu")
            pg.wait_for_timeout(300)
            checar(pg.evaluate("document.body.classList.contains('menu-aberto')"), "menu abre como gaveta")
            pg.click("#bt-novo")
            pg.fill("input[list=empresas-conhecidas]", "Salão Teste ME")
            pg.fill("input[inputmode=numeric]", "12345678000190")
            pg.fill("input[type=month]", "2026-09")
            pg.click("text=Criar documento")
            pg.wait_for_selector("#novo-texto")
            pg.fill("#novo-texto", "nf 321 sem boleto, valor 450,00")
            pg.select_option("#novo-prioridade", "Alta")
            pg.click("text=Adicionar apontamento")
            pg.fill("#novo-texto", "INVENTAR teste da conferência")
            pg.click("text=Adicionar apontamento")
            pg.wait_for_selector(".ap.formalizado >> nth=1", timeout=15000)
            checar(len(pedidos_ia) == 2, "IA chamada direto do aparelho")
            cab = pedidos_ia[0]["cab"]
            checar(cab.get("x-api-key") == "sk-ant-teste-aparelho" and
                   cab.get("anthropic-dangerous-direct-browser-access") == "true" and
                   cab.get("anthropic-version") == "2023-06-01", "cabeçalhos certos para a API")
            checar(pedidos_ia[0]["corpo"]["output_config"]["format"]["type"] == "json_schema",
                   "pede a resposta no formato estruturado")
            checar("Valor informado" not in pedidos_ia[0]["corpo"]["messages"][0]["content"],
                   "campo vazio não vai para a IA")
            avisos = pg.inner_text("#lista-ap")
            checar("20" in avisos and "não está no que você escreveu" in avisos, "conferência pega número inventado")
            checar(pg.locator(".coluna-previa").is_hidden(), "no celular a prévia fica em outra aba")
            pg.click(".abas-movel >> text=Prévia do documento")
            checar(pg.locator(".coluna-previa").is_visible() and "nf 321" in pg.inner_text(".papel"),
                   "prévia no celular")
            pg.click(".abas-movel >> text=Escrever")

            # finaliza: PDF e Word gerados no próprio aparelho
            pg.click("text=Finalizar e gerar")
            pg.wait_for_selector(".modal >> text=avisos de conferência")
            pg.click(".modal >> text=Gerar assim mesmo")
            pg.wait_for_selector(".faixa-final", timeout=30000)
            arquivos = {}
            for fmt, rot in (("pdf", "Baixar PDF"), ("docx", "Baixar Word")):
                with pg.expect_download() as baixa:
                    pg.click("text=" + rot)
                destino = TEMP / ("app." + fmt)
                baixa.value.save_as(destino)
                arquivos[fmt] = (destino, baixa.value.suggested_filename)
            from pypdf import PdfReader
            texto_pdf = "\n".join(pg_.extract_text() for pg_ in PdfReader(str(arquivos["pdf"][0])).pages)
            checar(arquivos["pdf"][1].endswith(".pdf") and "Salão Teste ME" in texto_pdf and
                   "PROVIDÊNCIA SOLICITADA" in texto_pdf and "Página 1 de" in texto_pdf, "PDF gerado no aparelho")
            from docx import Document
            dx = Document(str(arquivos["docx"][0]))
            tudo = "\n".join(p_.text for p_ in dx.paragraphs)
            checar("Identificamos o seguinte ponto: nf 321" in tudo and "Zera Contabilidade".upper() in tudo,
                   "Word gerado no aparelho")
            with zipfile.ZipFile(arquivos["docx"][0]) as z:
                checar(any(n.startswith("word/media/") for n in z.namelist()), "Word do aparelho com a logo")
            checar(not _ordem_xml_ok(arquivos["docx"][0]), "Word do aparelho com XML na ordem do padrão")
            for p_ in _ordem_xml_ok(arquivos["docx"][0])[:5]:
                print("        ", p_)

            pg.click("text=Mensagem para o cliente")
            checar("Salão Teste ME" in pg.input_value(".modal textarea"), "mensagem para o cliente")
            pg.click(".modal >> text=Fechar")

            # recarrega: os dados continuam no aparelho
            pg.reload()
            pg.wait_for_selector(".faixa-final")
            checar(pg.locator(".ap").count() == 2, "dados continuam no aparelho depois de fechar")

            # backup: baixa e importa em outro "aparelho" (outro perfil do navegador)
            pg.goto(base + "#/config")
            pg.wait_for_selector("text=Baixar backup")
            with pg.expect_download() as baixa:
                pg.click("text=Baixar backup")
            bk_caminho = TEMP / "backup.json"
            baixa.value.save_as(bk_caminho)
            bk = json.loads(bk_caminho.read_text(encoding="utf-8"))
            checar(bk["app"] == "apontamentos-contabeis-ia" and "sk-ant" not in bk_caminho.read_text(encoding="utf-8"),
                   "backup do aparelho sem a chave")
            # o mesmo backup entra no programa do Windows
            base_win = Base(TEMP / "windows.db")
            checar(base_win.importar(bk)["importados"] == 1, "backup do iPhone abre no programa do Windows")
            pg.goto(base + "#/")
            pg.wait_for_selector(".painel")
            pg.screenshot(path=str(TEMP / "celular.png"))

            ctx2 = nav.new_context(viewport={"width": 820, "height": 1180}, service_workers="block")
            pg2 = ctx2.new_page()
            pg2.on("pageerror", lambda e: erros.append(str(e)))
            pg2.on("dialog", lambda d: d.accept())
            pg2.goto(base + "#/config")
            pg2.wait_for_selector("text=Importar backup")
            pg2.set_input_files("input[type=file]", str(bk_caminho))
            pg2.wait_for_selector("text=1 documento(s) importado(s)")
            pg2.goto(base + "#/")
            pg2.wait_for_selector("text=Salão Teste ME")
            # trazer de outro mês
            pg2.goto(base + "#/novo")
            pg2.fill("input[list=empresas-conhecidas]", "Salão Teste ME")
            pg2.fill("input[type=month]", "2026-10")
            pg2.click("text=Criar documento")
            pg2.wait_for_selector(".modal >> text=Trazer marcados")
            pg2.wait_for_selector(".item-copiar >> nth=1")
            pg2.locator(".item-copiar input").nth(1).uncheck()
            pg2.click("text=Trazer marcados")
            pg2.wait_for_selector(".ap")
            checar(pg2.locator(".ap").count() == 1 and "Trazido de Setembro/2026" in pg2.inner_text("#lista-ap"),
                   "trazer pendência do mês anterior no aplicativo")
            ctx2.close()
            nav.close()
        checar(not erros, "sem erro de JavaScript no aplicativo")
        for e in erros[:5]:
            print("        ", e)
        print("   (captura do celular em", TEMP / "celular.png", ")")
    finally:
        srv.shutdown()
        srv.server_close()


_CLAUDE_FALSO = r"""
(() => {
  const CH = "claude-falso-db";
  const ler = () => { try { return JSON.parse(localStorage.getItem(CH) || "{}"); } catch (e) { return {}; } };
  const gravar = (m) => localStorage.setItem(CH, JSON.stringify(m));
  const espera = () => new Promise(r => setTimeout(r, 5));
  const snap = (path, v) => ({ id: path.split("/").pop(), exists: v !== undefined,
                               data: () => (v === undefined ? undefined : Object.freeze(JSON.parse(JSON.stringify(v)))),
                               metadata: { fromCache: false, hasPendingWrites: false } });
  function checar(path, par) {
    const n = path.split("/").length;
    if ((n % 2 === 0) !== par) throw new TypeError("caminho com paridade errada: " + path);
  }
  function doc(path) {
    checar(path, true);
    return { id: path.split("/").pop(), path,
      get: async () => { await espera(); return snap(path, ler()[path]); },
      set: async (d) => { await espera(); if (JSON.stringify(d).length > 262144) throw { code: "invalid_argument" };
                          const m = ler(); m[path] = JSON.parse(JSON.stringify(d)); gravar(m); },
      update: async (d) => { const m = ler(); Object.assign(m[path], d); gravar(m); },
      delete: async () => { await espera(); const m = ler(); delete m[path]; gravar(m); },
      collection: (sub) => colecao(path + "/" + sub) };
  }
  function colecao(path, filtros, lim) {
    checar(path, false);
    filtros = filtros || [];
    return { path,
      where: (f, op, v) => colecao(path, filtros.concat([[f, op, v]]), lim),
      orderBy: () => colecao(path, filtros, lim),
      limit: (n) => colecao(path, filtros, n),
      doc: (id) => doc(path + "/" + (id || Math.random().toString(36).slice(2))),
      get: async () => {
        await espera();
        const m = ler();
        const docs = Object.keys(m).filter(k => k.startsWith(path + "/") && k.split("/").length === path.split("/").length + 1)
          .filter(k => filtros.every(([f, op, v]) => op === "==" && m[k][f] === v))
          .slice(0, lim || 1000).map(k => snap(k, m[k]));
        return { docs, size: docs.length, empty: !docs.length, docChanges: () => [], metadata: {} };
      } };
  }
  const db = { doc, collection: (p) => colecao(p) };
  window.__baixados = [];
  window.__pedidosIA = [];
  const sample = async () => { throw { code: "invalid_request" }; };
  sample.json = async (pedido) => {
    window.__pedidosIA.push(pedido);
    await new Promise(r => setTimeout(r, 300));
    const anot = pedido.split("Anotação do contador:\n").pop();
    const r = { titulo: "Apontamento formalizado", texto: "Identificamos o seguinte ponto: " + anot + ".",
                providencia: "Solicitamos a verificação do item.", pontos_a_confirmar: [] };
    if (anot.includes("INVENTAR")) r.texto += " Multa de 20%.";
    return r;
  };
  const downloads = { save: async ({ filename, data }) => {
    const buf = new Uint8Array(await data.arrayBuffer());
    let bin = ""; for (let i = 0; i < buf.length; i += 0x8000) bin += String.fromCharCode.apply(null, buf.subarray(i, i + 0x8000));
    window.__baixados.push({ filename, b64: btoa(bin) });
    return { status: "saved" };
  } };
  const user = { id: async () => "u1", isOwner: () => true };
  const caps = { db, sample, downloads, user };
  window.claude = { use: async (n) => { await espera(); return caps[n] || null; } };
})();
"""


class _ServidorClaude(BaseHTTPRequestHandler):
    """Faz o papel do claude.ai: serve a página montada, com o esqueleto que a publicação acrescenta."""
    def log_message(self, *a):
        pass

    def do_GET(self):
        from urllib.parse import unquote, urlparse
        caminho = unquote(urlparse(self.path).path).lstrip("/") or "index.html"
        alvo = (RAIZ / "claude" / caminho).resolve()
        if not alvo.is_file() or (RAIZ / "claude") not in alvo.parents:
            self.send_response(404)
            self.end_headers()
            return
        dados = alvo.read_bytes()
        tipo = {".html": "text/html; charset=utf-8", ".png": "image/png", ".ttf": "font/ttf"}.get(alvo.suffix, "application/octet-stream")
        if alvo.name == "index.html":
            dados = (b'<!doctype html><html><head><meta charset=utf8><meta name=viewport '
                     b'content="width=device-width,initial-scale=1,viewport-fit=cover"></head><body>'
                     + dados + b"</body></html>")
        self.send_response(200)
        self.send_header("Content-Type", tipo)
        self.send_header("Content-Length", str(len(dados)))
        self.end_headers()
        self.wfile.write(dados)


def teste_claude():
    """A versão que abre dentro do Claude: sem chave, dados na conta, IA da conta."""
    print("\n[versão dentro do Claude (link do claude.ai)]")
    try:
        from playwright.sync_api import sync_playwright
    except ImportError:
        print("  (pulado: playwright não instalado)")
        return
    sys.path.insert(0, str(RAIZ / "ferramentas"))
    import montar_claude
    pagina = montar_claude.montar()
    checar(pagina.stat().st_size < 16_000_000, "página dentro do limite de tamanho")
    texto = pagina.read_text(encoding="utf-8")
    checar("<html" not in texto[:2000] and "<title>" in texto[:200], "página no formato da publicação")
    from http.server import ThreadingHTTPServer
    srv = ThreadingHTTPServer(("127.0.0.1", 0), _ServidorClaude)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    base = f"http://127.0.0.1:{srv.server_port}/"
    erros = []
    try:
        with sync_playwright() as p:
            caminho = os.environ.get("CHROMIUM") or ("/opt/pw-browsers/chromium-1194/chrome-linux/chrome"
                                                    if Path("/opt/pw-browsers/chromium-1194/chrome-linux/chrome").exists() else None)
            nav = p.chromium.launch(executable_path=caminho) if caminho else p.chromium.launch()
            ctx = nav.new_context(viewport={"width": 390, "height": 844}, is_mobile=True, has_touch=True,
                                  service_workers="block")
            ctx.add_init_script(_CLAUDE_FALSO)
            ctx.route("https://api.anthropic.com/**", lambda r: (erros.append("chamou a API direto"), r.abort()))
            pg = ctx.new_page()
            pg.on("pageerror", lambda e: erros.append(str(e)))
            pg.on("console", lambda m: erros.append(m.text) if m.type == "error" and "404" not in m.text else None)
            pg.on("dialog", lambda d: (erros.append("usou diálogo do navegador: " + d.message), d.dismiss()))
            pg.goto(base)
            pg.wait_for_selector(".painel")
            checar("no Claude" in pg.inner_text("#versao"), "abre no modo Claude")
            checar("IA da sua conta Claude" in pg.inner_text("#selo-ia") or pg.locator("text=configurar agora").count() == 0,
                   "não pede chave da API")
            checar(pg.locator("text=Configurar agora").count() == 0, "sem aviso de chave")
            pg.click("#bt-config")
            pg.wait_for_selector("text=A IA usa a sua própria conta do Claude")
            checar(pg.locator("input[type=password]").count() == 0, "configurações sem campo de chave")
            pg.locator("input.campo").nth(0).fill("Zera Contabilidade")
            pg.click("text=Salvar configurações")
            pg.wait_for_selector("text=Configurações salvas.")
            pg.goto(base + "#/novo")
            pg.fill("input[list=empresas-conhecidas]", "Clínica Teste Ltda")
            pg.fill("input[type=month]", "2026-09")
            pg.click("text=Criar documento")
            pg.wait_for_selector("#novo-texto")
            checar(pg.locator(".bt-mic").count() == 0, "sem ditado (microfone bloqueado no Claude)")
            pg.fill("#novo-texto", "nf 777 sem boleto, valor 300,00")
            pg.click("text=Adicionar apontamento")
            pg.fill("#novo-texto", "INVENTAR item para excluir")
            pg.click("text=Adicionar apontamento")
            pg.wait_for_selector(".ap.formalizado >> nth=1", timeout=15000)
            pedidos = pg.evaluate("window.__pedidosIA")
            checar(len(pedidos) == 2 and "Não invente nada" in pedidos[0] and "Anotação do contador:\nnf 777" in pedidos[0],
                   "IA da conta recebe as mesmas instruções")
            checar("não está no que você escreveu" in pg.inner_text("#lista-ap"), "conferência de números no Claude")
            pg.locator(".ap").nth(1).locator("button:has-text('Excluir')").click()
            pg.wait_for_selector(".modal >> text=Excluir o apontamento 2?")
            pg.click(".modal >> button:has-text('Excluir')")
            pg.wait_for_function("document.querySelectorAll('.ap').length === 1")
            checar(True, "excluir com confirmação dentro da página")
            chaves = pg.evaluate("Object.keys(JSON.parse(localStorage.getItem('claude-falso-db')))")
            checar(chaves and all(k.startswith("data/users/u1/apontamentos/") for k in chaves),
                   "dados guardados só na área do próprio usuário")
            pg.click("text=Finalizar e gerar")
            pg.wait_for_selector(".faixa-final", timeout=30000)
            checar(pg.locator("text=Compartilhar PDF").count() == 0, "sem botão de compartilhar (bloqueado no Claude)")
            pg.click("text=Baixar PDF")
            pg.wait_for_function("window.__baixados.length === 1")
            import base64
            from pypdf import PdfReader
            b = pg.evaluate("window.__baixados[0]")
            (TEMP / "claude.pdf").write_bytes(base64.b64decode(b["b64"]))
            texto_pdf = "\n".join(x.extract_text() for x in PdfReader(str(TEMP / "claude.pdf")).pages)
            checar(b["filename"].endswith(".pdf") and "Clínica Teste Ltda" in texto_pdf, "PDF salvo pelo Claude")
            # reabre a página: os documentos continuam (na conta) e o arquivo é gerado de novo
            pg.goto(base + "#/")
            pg.reload()
            pg.wait_for_selector(".painel")
            checar("Clínica Teste Ltda" in pg.inner_text(".painel"), "documentos continuam ao reabrir")
            pg.locator(".doc-recente").first.click()
            pg.wait_for_selector(".faixa-final")
            pg.click("text=Baixar Word")
            pg.wait_for_function("window.__baixados.length === 1")
            b = pg.evaluate("window.__baixados[0]")
            (TEMP / "claude.docx").write_bytes(base64.b64decode(b["b64"]))
            from docx import Document
            checar("Identificamos o seguinte ponto: nf 777" in "\n".join(x.text for x in Document(str(TEMP / "claude.docx")).paragraphs),
                   "Word gerado de novo depois de reabrir")
            pg.screenshot(path=str(TEMP / "claude.png"))
            nav.close()
        checar(not erros, "sem erro de JavaScript na versão do Claude")
        for e in erros[:5]:
            print("        ", e)
    finally:
        srv.shutdown()
        srv.server_close()


def main():
    for teste in (teste_conferencia, teste_redator, teste_pedido_real_da_biblioteca,
                  teste_config, teste_aplicacao, teste_servidor, teste_tela,
                  teste_paridade_js, teste_aplicativo, teste_claude):
        try:
            teste()
        except Exception:
            traceback.print_exc()
            FALHAS.append(teste.__name__ + " (exceção)")
    print()
    if FALHAS:
        print(f"RESULTADO: FALHOU ({len(FALHAS)})")
        for f in FALHAS:
            print("  -", f)
        return 1
    print("RESULTADO: OK")
    return 0


if __name__ == "__main__":
    sys.exit(main())
